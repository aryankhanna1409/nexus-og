#!/usr/bin/env python3
"""Builds AegisNet_Technical_Handbook.docx from deliverables/documentation/handbook_source.md.

The Markdown file is the canonical, editable source. This script maps a small Markdown dialect to
native Word structures (styles, tables, fields, captions, bookmarks); nothing is rasterised except
the figures themselves. Dialect:

  # Title                  chapter (numbered automatically; "Appendix ..." titles are not numbered)
  ## / ###                 numbered sections
  - item / 1. item         lists
  Table: caption           caption for the following pipe table
  | a | b |                tables (first row is the header)
  Figure: caption {id=x width=6.5}  + next line ![alt](path)     figure with caption; {@x} cites it
  > [!NOTE|WARNING|LIMITATION|EXAMPLE|STATUS] text               callout box
  ```lang ... ```          code block;   $$ formula $$ centred formula
  **bold** *italic* `code` [[V]] [[U]] [[S]] [[N]] status badges
"""

from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor
from PIL import Image

REPO = Path(__file__).resolve().parents[2]
DELIV = REPO / "deliverables"
SOURCE = DELIV / "documentation" / "handbook_source.md"
OUT = DELIV / "handbook" / "AegisNet_Technical_Handbook.docx"

NAVY, BLUE, INK, MUTED = "0B2545", "1D4ED8", "1B2A41", "5B6B7F"
FONT, MONO, MATH = "Calibri", "Consolas", "Cambria Math"
CONTENT_CM = 17.0

CALLOUTS = {  # label -> (border, fill)
    "NOTE": ("1D4ED8", "EFF6FF"), "WARNING": ("B91C1C", "FEF2F2"), "LIMITATION": ("B45309", "FFFBEB"),
    "EXAMPLE": ("0F766E", "F0FDFA"), "STATUS": ("475569", "F1F5F9"),
}
BADGES = {  # tag -> (label, text colour, fill)
    "V": ("VERIFIED", "14532D", "DCFCE7"), "U": ("UNVERIFIED HERE", "7C2D12", "FFEDD5"),
    "S": ("SIMULATED", "4C1D95", "EDE9FE"), "N": ("NOT IMPLEMENTED", "7F1D1D", "FEE2E2"),
}


def git_head() -> str:
    import subprocess
    try:
        return subprocess.run(["git", "-C", str(REPO), "rev-parse", "--short", "HEAD"], capture_output=True,
                              text=True, check=True).stdout.strip()
    except Exception:  # noqa: BLE001
        return "unknown"


# ---- low-level XML helpers ----------------------------------------------------------------------


def set_shading(element, fill: str) -> None:  # noqa: ANN001
    props = element.get_or_add_tcPr() if hasattr(element, "get_or_add_tcPr") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def cell_borders(cell, **edges: tuple[str, int] | None) -> None:  # noqa: ANN001
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        spec = edges.get(edge)
        el = OxmlElement(f"w:{edge}")
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:color"), spec[0])
            el.set(qn("w:sz"), str(spec[1]))
            el.set(qn("w:space"), "0")
        borders.append(el)
    tcPr.append(borders)


def cell_margins(cell, top=60, bottom=60, left=100, right=100) -> None:  # noqa: ANN001
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def add_field(paragraph, instr: str, placeholder: str = "", bold: bool = False, size: float | None = None):  # noqa: ANN001
    """Inserts a complex field (PAGE, SEQ, TOC, REF...) with cached placeholder text."""
    def run_with(child):  # noqa: ANN001, ANN202
        r = paragraph.add_run()
        if bold:
            r.bold = True
        if size:
            r.font.size = Pt(size)
        r._r.append(child)
        return r

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run_with(begin)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run_with(it)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run_with(sep)
    r = paragraph.add_run(placeholder)
    if bold:
        r.bold = True
    if size:
        r.font.size = Pt(size)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)


_bookmark_id = [100]


def bookmark(paragraph, name: str, runs_start: int) -> None:  # noqa: ANN001
    """Wraps the paragraph's runs from index `runs_start` in a named bookmark."""
    _bookmark_id[0] += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(_bookmark_id[0]))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(_bookmark_id[0]))
    runs = paragraph._p.findall(qn("w:r"))
    runs[runs_start].addprevious(start)
    runs[-1].addnext(end)


# ---- document setup -------------------------------------------------------------------------------


def configure_styles(doc: Document) -> None:
    st = doc.styles
    normal = st["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    def heading(name: str, size: float, color: str, before: float, after: float, page_break: bool = False) -> None:
        s = st[name]
        s.font.name, s.font.size, s.font.bold = FONT, Pt(size), True
        s.element.rPr.rFonts.set(qn("w:ascii"), FONT)
        s.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        s.font.color.rgb = RGBColor.from_string(color)
        pf = s.paragraph_format
        pf.space_before, pf.space_after, pf.keep_with_next = Pt(before), Pt(after), True
        pf.page_break_before = page_break

    heading("Heading 1", 24, NAVY, 0, 14, page_break=True)
    heading("Heading 2", 15, NAVY, 16, 6)
    heading("Heading 3", 12, BLUE, 12, 4)

    for name, size, color in (("Title", 34, NAVY),):
        s = st[name]
        s.font.name, s.font.size, s.font.bold = FONT, Pt(size), True
        s.font.color.rgb = RGBColor.from_string(color)

    cap = st["Caption"]
    cap.font.name, cap.font.size, cap.font.bold, cap.font.italic = FONT, Pt(9.5), True, False
    cap.font.color.rgb = RGBColor.from_string(MUTED)
    cap.paragraph_format.space_before, cap.paragraph_format.space_after = Pt(3), Pt(10)

    for lvl, left in ((1, 0.0), (2, 0.5)):
        name = f"TOC {lvl}"
        s = st[name] if name in [x.name for x in st] else st.add_style(name, 1)
        s.font.name, s.font.size = FONT, Pt(10.5 if lvl == 1 else 10)
        s.font.bold = lvl == 1
        s.paragraph_format.left_indent = Cm(left)
        s.paragraph_format.space_after = Pt(2)

    code = st.add_style("CodeBlock", 1)
    code.font.name, code.font.size = MONO, Pt(8.3)
    code.element.rPr.rFonts.set(qn("w:eastAsia"), MONO)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.left_indent = Cm(0.2)

    tb = st.add_style("TableText", 1)
    tb.font.name, tb.font.size = FONT, Pt(8.8)
    tb.paragraph_format.space_after = Pt(0)
    tb.paragraph_format.line_spacing = 1.05

    bl = st["List Bullet"]
    bl.font.name, bl.font.size = FONT, Pt(10)
    bl.paragraph_format.space_after = Pt(3)
    nl = st["List Number"]
    nl.font.name, nl.font.size = FONT, Pt(10)
    nl.paragraph_format.space_after = Pt(3)


def page_setup(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.2)
    sec.header_distance, sec.footer_distance = Cm(1.0), Cm(0.9)
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.text = ""
    r = hp.add_run("Aegis-Net Technical Handbook")
    r.font.size, r.font.color.rgb = Pt(8.5), RGBColor.from_string(MUTED)
    hp.add_run("\t")
    r2 = hp.add_run(f"Version 1.0 - repository commit {git_head()}")
    r2.font.size, r2.font.color.rgb = Pt(8.5), RGBColor.from_string(MUTED)
    pPr = hp._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "4"), ("space", "4"), ("color", "CBD5E1")):
        b.set(qn(f"w:{k}"), v)
    borders.append(b)
    pPr.append(borders)
    tabs = OxmlElement("w:tabs")
    for pos, kind in ((9638, "right"),):
        t = OxmlElement("w:tab")
        t.set(qn("w:val"), kind)
        t.set(qn("w:pos"), str(pos))
        tabs.append(t)
    pPr.append(tabs)
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    r = fp.add_run("All data is synthetic. Aegis-Net is a prototype, not a production system.")
    r.font.size, r.font.color.rgb = Pt(8.5), RGBColor.from_string(MUTED)
    fp.add_run("\t")
    pr = fp.add_run("Page ")
    pr.font.size, pr.font.color.rgb = Pt(8.5), RGBColor.from_string(MUTED)
    add_field(fp, "PAGE", "1", size=8.5)
    fpPr = fp._p.get_or_add_pPr()
    tabs2 = OxmlElement("w:tabs")
    for pos, kind in ((9638, "right"),):
        t = OxmlElement("w:tab")
        t.set(qn("w:val"), kind)
        t.set(qn("w:pos"), str(pos))
        tabs2.append(t)
    fpPr.append(tabs2)


def cover_page(doc: Document) -> None:
    def para(text: str, size: float, color: str = INK, bold: bool = False, after: float = 6, before: float = 0,
             italic: bool = False):  # noqa: ANN202
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.size, r.bold, r.italic = Pt(size), bold, italic
        r.font.color.rgb = RGBColor.from_string(color)
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        return p

    para("", 10, after=90)
    bar = doc.add_paragraph()
    pPr = bar._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    b = OxmlElement("w:top")
    for k, v in (("val", "single"), ("sz", "36"), ("space", "6"), ("color", BLUE)):
        b.set(qn(f"w:{k}"), v)
    borders.append(b)
    pPr.append(borders)
    para("AEGIS-NET", 14, BLUE, bold=True, after=4, before=6)
    p = doc.add_paragraph(style="Title")
    p.add_run("Technical Handbook")
    p.paragraph_format.space_after = Pt(4)
    para("Network intelligence, AI root-cause analysis and a policy-gated DevSecOps agent: from first principles to "
         "implementation details", 15, MUTED, after=28)
    para("A learning guide, architecture reference and project handover for the Aegis-Net prototype.", 11.5, INK, after=60)
    meta = [("Version", "1.0"), ("Date", date.today().isoformat()),
            ("Repository state", f"commit {git_head()} on branch master (working tree includes the documentation under "
                                 "deliverables/)"),
            ("Author field", "Aryan Khanna"),
            ("Data notice", "All data in Aegis-Net is synthetic. No claim is made about real networks.")]
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for k, v in meta:
        row = t.add_row()
        row.cells[0].width, row.cells[1].width = Cm(3.6), Cm(12.8)
        for c, text, bold in ((row.cells[0], k, True), (row.cells[1], v, False)):
            c.text = ""
            r = c.paragraphs[0].add_run(text)
            r.bold, r.font.size = bold, Pt(10)
            r.font.color.rgb = RGBColor.from_string(MUTED if bold else INK)
            cell_borders(c, bottom=("E2E8F0", 4))
            cell_margins(c, 70, 70, 60, 60)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---- inline markup -------------------------------------------------------------------------------

TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[\[[VUSN]\]\]|\{@[A-Za-z0-9_-]+\}|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))")


class Builder:
    def __init__(self, doc: Document) -> None:
        self.doc = doc
        self.fig_no = 0
        self.tab_no = 0
        self.fig_ids: dict[str, int] = {}
        self.chapter = 0
        self.section = 0
        self.subsection = 0
        self.pending_table_caption: str | None = None
        self.pending_figure: tuple[str, str, float] | None = None

    # --- inline ---
    def inline(self, p, text: str, size: float | None = None, base_bold: bool = False, color: str | None = None) -> None:  # noqa: ANN001
        pos = 0
        for m in TOKEN.finditer(text):
            if m.start() > pos:
                self._run(p, text[pos:m.start()], size, base_bold, color)
            tok = m.group(0)
            if tok.startswith("**"):
                self._run(p, tok[2:-2], size, True, color)
            elif tok.startswith("`"):
                code = tok[1:-1]
                if re.search(r"\d x |\bx \d| x max| x \(", code):
                    code = code.replace(" x ", " × ")
                r = self._run(p, code, (size or 10.5) - 1, base_bold, "1E3A8A")
                r.font.name = MONO
                r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), MONO)
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "EEF2F7")
                r._r.get_or_add_rPr().append(shd)
            elif tok.startswith("[["):
                label, fg, bg = BADGES[tok[2]]
                r = self._run(p, f" {label} ", (size or 10.5) - 2.2, True, fg)
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), bg)
                r._r.get_or_add_rPr().append(shd)
            elif tok.startswith("{@"):
                self._ref(p, tok[2:-1], size)
            else:
                r = self._run(p, tok[1:-1], size, base_bold, color)
                r.italic = True
            pos = m.end()
        if pos < len(text):
            self._run(p, text[pos:], size, base_bold, color)

    def _run(self, p, text: str, size, bold: bool, color):  # noqa: ANN001, ANN202
        r = p.add_run(text)
        if size:
            r.font.size = Pt(size)
        if bold:
            r.bold = True
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return r

    def _ref(self, p, ident: str, size) -> None:  # noqa: ANN001
        number = self.fig_ids.get(ident)
        text = f"Figure {number}" if number else f"Figure ?{ident}"
        add_field(p, f"REF fig_{ident.replace('-', '_')} \\h", text, size=size)

    # --- blocks ---
    def heading(self, level: int, text: str) -> None:
        if level == 1:
            if text.startswith("Appendix"):
                label = text
            else:
                self.chapter += 1
                self.section = self.subsection = 0
                label = f"{self.chapter}  {text}"
        elif level == 2:
            self.section += 1
            self.subsection = 0
            label = f"{self.chapter}.{self.section}  {text}" if self.chapter else text
        else:
            self.subsection += 1
            label = f"{self.chapter}.{self.section}.{self.subsection}  {text}"
        hp = self.doc.add_heading(label, level=min(level, 3))
        if level == 1 and text.startswith(("Appendix B", "Appendix C")):
            hp.paragraph_format.page_break_before = False
            hp.paragraph_format.space_before = Pt(24)

    def paragraph(self, text: str) -> None:
        p = self.doc.add_paragraph()
        self.inline(p, text)

    def bullets(self, items: list[str], numbered: bool) -> None:
        style = "List Number" if numbered else "List Bullet"
        for it in items:
            p = self.doc.add_paragraph(style=style)
            self.inline(p, it)
        if numbered:
            self._restart_numbering(items)

    def _restart_numbering(self, items) -> None:  # noqa: ANN001
        pass  # List Number continues across lists; lists in this handbook are short and rendered via text numbers

    def numbered_as_text(self, items: list[str]) -> None:
        for i, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(3)
            self._run(p, f"{i}.  ", None, True, BLUE)
            self.inline(p, it)

    def code(self, lines: list[str]) -> None:
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]
        c.width = Cm(CONTENT_CM)
        set_shading(c._tc, "F4F6F9")
        cell_borders(c, left=("1D4ED8", 18), top=("E2E8F0", 4), bottom=("E2E8F0", 4), right=("E2E8F0", 4))
        cell_margins(c, 80, 80, 140, 100)
        c.paragraphs[0].style = self.doc.styles["CodeBlock"]
        first = True
        for line in lines:
            p = c.paragraphs[0] if first else c.add_paragraph(style="CodeBlock")
            first = False
            r = p.add_run(line if line else " ")
            r.font.name = MONO
            r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), MONO)
            p.paragraph_format.keep_together = True
        self._cant_split(t.rows[0])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    @staticmethod
    def _cant_split(row) -> None:  # noqa: ANN001
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:cantSplit")
        trPr.append(el)

    def formula(self, text: str) -> None:
        text = text.replace(" x ", " × ")
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F8FAFC")
        pPr.append(shd)
        r = p.add_run(text)
        r.font.name, r.font.size, r.italic = MATH, Pt(11.5), True
        r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), MATH)
        r.font.color.rgb = RGBColor.from_string(NAVY)

    def callout(self, kind: str, text: str) -> None:
        border, fill = CALLOUTS[kind]
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]
        c.width = Cm(CONTENT_CM)
        set_shading(c._tc, fill)
        cell_borders(c, left=(border, 24))
        cell_margins(c, 90, 90, 160, 140)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(kind.title() + ".  " if kind != "STATUS" else "Implementation status.  ")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(border)
        self.inline(p, text, size=9.5)
        self._cant_split(t.rows[0])
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(3)

    def table(self, rows: list[list[str]]) -> None:
        caption = self.pending_table_caption
        self.pending_table_caption = None
        if caption:
            self.tab_no += 1
            cp = self.doc.add_paragraph(style="Caption")
            cp.paragraph_format.keep_with_next = True
            cp.paragraph_format.space_after = Pt(3)
            cp.add_run("Table ")
            add_field(cp, "SEQ Table \\* ARABIC", str(self.tab_no), bold=True)
            cp.add_run(f". {caption}")
        ncols = len(rows[0])
        plain = lambda v: re.sub(r"[`*{}@]|\[\[|\]\]", "", v)  # noqa: E731
        weights, mins = [], []
        for c in range(ncols):
            cells = [plain(r[c]) if c < len(r) else "" for r in rows]
            longest = max(len(x) for x in cells)
            avg = sum(len(x) for x in cells) / len(cells)
            weights.append(min(max(0.6 * avg + 0.4 * longest, 6), 70) ** 0.9)
            word = max((len(w) for x in cells for w in x.split()), default=4)
            word = max(word, 15 if any(r[c].count("[[") for r in rows if c < len(r)) else 0)  # badge labels
            has_code = any("`" in r[c] for r in rows if c < len(r))
            mins.append(max(1.0, min(word * (0.2 if has_code else 0.175) + 0.5, 6.0)))
        if sum(mins) > CONTENT_CM * 0.8:  # never reserve more than 80 % of the width for minimums
            mins = [m * CONTENT_CM * 0.8 / sum(mins) for m in mins]
        total = sum(weights)
        cm = [CONTENT_CM * w / total for w in weights]
        for _ in range(4):  # lift narrow columns to their minimum, take the space from wide ones
            short = [i for i in range(ncols) if cm[i] < mins[i]]
            if not short:
                break
            need = min(sum(mins[i] - cm[i] for i in short), sum(max(cm[i] - mins[i], 0) for i in range(ncols) if i not in short) * 0.9)
            rich = [i for i in range(ncols) if cm[i] > mins[i] + 0.5 and i not in short]
            room = sum(cm[i] - mins[i] for i in rich)
            if room <= 0:
                break
            scale_short = need / max(sum(mins[i] - cm[i] for i in short), 1e-9)
            for i in short:
                cm[i] += (mins[i] - cm[i]) * scale_short
            for i in rich:
                cm[i] -= need * (cm[i] - mins[i]) / room
        widths = [Cm(v) for v in cm]
        t = self.doc.add_table(rows=len(rows), cols=ncols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        for ri, row in enumerate(rows):
            tr = t.rows[ri]
            self._cant_split(tr)
            for ci in range(ncols):
                cell = tr.cells[ci]
                cell.width = widths[ci]
                cell.text = ""
                p = cell.paragraphs[0]
                p.style = self.doc.styles["TableText"]
                text = row[ci] if ci < len(row) else ""
                if ri == 0:
                    set_shading(cell._tc, NAVY)
                    self.inline(p, text, size=8.8, base_bold=True, color="FFFFFF")
                else:
                    if ri % 2 == 0:
                        set_shading(cell._tc, "F5F8FC")
                    self.inline(p, text, size=8.8)
                cell_borders(cell, top=("D5DEE9", 4), bottom=("D5DEE9", 4), left=("D5DEE9", 4), right=("D5DEE9", 4))
                cell_margins(cell, 50, 50, 80, 80)
            if ri == 0:
                trPr = tr._tr.get_or_add_trPr()
                h = OxmlElement("w:tblHeader")
                trPr.append(h)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)

    def figure(self, caption: str, ident: str, width_in: float, path: str) -> None:
        img = DELIV / path
        if not img.is_file():
            raise FileNotFoundError(img)
        with Image.open(img) as im:
            w, h = im.size
        width = min(width_in, CONTENT_CM / 2.54)
        height = width * h / w
        max_h = 8.1
        if height > max_h:
            width, height = width * max_h / height, max_h
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(img), width=Emu(int(width * 914400)))
        self.fig_no += 1
        cp = self.doc.add_paragraph(style="Caption")
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        start = len(cp._p.findall(qn("w:r")))
        cp.add_run("Figure ")
        add_field(cp, "SEQ Figure \\* ARABIC", str(self.fig_no), bold=True)
        bookmark(cp, f"fig_{ident.replace('-', '_')}", start)
        cp.add_run(f". {caption}")
        # alt text
        for dp in p._p.iter(qn("wp:docPr")):
            dp.set("descr", f"Figure {self.fig_no}: {caption}")
            dp.set("title", f"Figure {self.fig_no}")


def parse_blocks(md: str) -> list[tuple]:
    lines = md.split("\n")
    blocks: list[tuple] = []
    i = 0
    para: list[str] = []

    def flush() -> None:
        nonlocal para
        if para:
            blocks.append(("p", " ".join(s.strip() for s in para)))
            para = []

    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            flush()
            j = i + 1
            code: list[str] = []
            while not lines[j].startswith("```"):
                code.append(lines[j])
                j += 1
            blocks.append(("code", code))
            i = j + 1
            continue
        m = re.match(r"^(#{1,3}) (.+)$", line)
        if m:
            flush()
            blocks.append(("h", len(m.group(1)), m.group(2).strip()))
        elif line.startswith("$$") and line.endswith("$$") and len(line) > 4:
            flush()
            blocks.append(("formula", line[2:-2].strip()))
        elif line.startswith("> [!"):
            flush()
            kind = re.match(r"> \[!(\w+)\]", line).group(1)
            text = re.sub(r"^> \[!\w+\]\s*", "", line)
            j = i + 1
            while j < len(lines) and lines[j].startswith("> ") and not lines[j].startswith("> [!"):
                text += " " + lines[j][2:].strip()
                j += 1
            blocks.append(("callout", kind, text))
            i = j
            continue
        elif line.startswith("Table: "):
            flush()
            blocks.append(("tcap", line[7:].strip()))
        elif line.startswith("Figure: "):
            flush()
            cap = line[8:].strip()
            m2 = re.search(r"\{id=([\w-]+) width=([\d.]+)\}\s*$", cap)
            ident, width = m2.group(1), float(m2.group(2))
            cap = cap[:m2.start()].strip()
            img = re.match(r"!\[.*?\]\((.+?)\)", lines[i + 1]).group(1)
            blocks.append(("fig", cap, ident, width, img))
            i += 2
            continue
        elif line.startswith("|"):
            flush()
            rows = []
            j = i
            while j < len(lines) and lines[j].startswith("|"):
                cells = [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", lines[j].strip()[1:-1])]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                j += 1
            blocks.append(("table", rows))
            i = j
            continue
        elif re.match(r"^- ", line):
            flush()
            items = []
            j = i
            while j < len(lines) and re.match(r"^- ", lines[j]):
                items.append(lines[j][2:].strip())
                j += 1
            blocks.append(("ul", items))
            i = j
            continue
        elif re.match(r"^\d+\. ", line):
            flush()
            items = []
            j = i
            while j < len(lines) and re.match(r"^\d+\. ", lines[j]):
                items.append(re.sub(r"^\d+\. ", "", lines[j]).strip())
                j += 1
            blocks.append(("ol", items))
            i = j
            continue
        elif not line.strip():
            flush()
        else:
            para.append(line)
        i += 1
    flush()
    return blocks


def build() -> Path:
    md = SOURCE.read_text(encoding="utf-8")
    blocks = parse_blocks(md)
    doc = Document()
    configure_styles(doc)
    page_setup(doc)
    cover_page(doc)
    b = Builder(doc)
    # first pass: figure numbers (so {@id} can show the right number before the figure appears)
    n = 0
    for blk in blocks:
        if blk[0] == "fig":
            n += 1
            b.fig_ids[blk[2]] = n
    # table of contents and list of figures
    h = doc.add_paragraph()
    r = h.add_run("Contents")
    r.bold, r.font.size = True, Pt(22)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-2" \\h \\z \\u', "Open in Word and press F9 (or right-click, Update Field) to build the table of contents.")
    h2 = doc.add_paragraph()
    h2.paragraph_format.page_break_before = True
    r = h2.add_run("List of figures")
    r.bold, r.font.size = True, Pt(18)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    lof = doc.add_paragraph()
    add_field(lof, 'TOC \\h \\z \\c "Figure"', "Update this field in Word to list the figures.")
    for blk in blocks:
        kind = blk[0]
        if kind == "h":
            b.heading(blk[1], blk[2])
        elif kind == "p":
            b.paragraph(blk[1])
        elif kind == "ul":
            b.bullets(blk[1], False)
        elif kind == "ol":
            b.numbered_as_text(blk[1])
        elif kind == "code":
            b.code(blk[1])
        elif kind == "formula":
            b.formula(blk[1])
        elif kind == "callout":
            b.callout(blk[1], blk[2])
        elif kind == "tcap":
            b.pending_table_caption = blk[1]
        elif kind == "table":
            b.table(blk[1])
        elif kind == "fig":
            b.figure(blk[1], blk[2], blk[3], blk[4])
    cp = doc.core_properties
    cp.title = "Aegis-Net Technical Handbook"
    cp.subject = "Architecture, implementation and operation of the Aegis-Net prototype"
    cp.author = "Aryan Khanna"
    cp.keywords = "Aegis-Net, anomaly detection, root-cause analysis, DevSecOps"
    cp.comments = f"Generated from deliverables/documentation/handbook_source.md at commit {git_head()}"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT.relative_to(REPO)} ({b.fig_no} figures, {b.tab_no} captioned tables)")
    return OUT


if __name__ == "__main__":
    sys.exit(0 if build() else 1)

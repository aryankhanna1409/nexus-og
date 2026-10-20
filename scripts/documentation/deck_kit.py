"""Helpers for building the Aegis-Net deck with native, editable PowerPoint objects."""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY, BLUE, INK, MUTED, LINE = "0B2545", "1D4ED8", "1B2A41", "5B6B7F", "8A99AD"
KINDS = {
    "native": ("FEF3C7", "B45309"), "py": ("DBEAFE", "1D4ED8"), "ai": ("EDE9FE", "6D28D9"),
    "store": ("CCFBF1", "0F766E"), "ui": ("DCFCE7", "15803D"), "sec": ("FEE2E2", "B91C1C"),
    "ext": ("F1F5F9", "64748B"), "plain": ("FFFFFF", "64748B"),
}
FONT = "Calibri"
SW, SH = 13.333, 7.5


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    return prs


def blank(prs: Presentation):  # noqa: ANN201
    return prs.slides.add_slide(prs.slide_layouts[6])


def text(slide, x: float, y: float, w: float, h: float, content, size: float = 14, color: str = INK, bold: bool = False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic: bool = False, font: str = FONT, line_spacing: float = 1.05,
         name: str | None = None):  # noqa: ANN001, ANN201
    """Text box. `content` is a string or a list of paragraphs; a paragraph is a string or a list of
    (text, {bold, color, size, italic}) runs. Strings starting with '- ' become bullets."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        runs = para if isinstance(para, list) else [(para, {})]
        bullet = False
        if runs and isinstance(runs[0][0], str) and runs[0][0].startswith("- "):
            bullet = True
            runs = [(runs[0][0][2:], runs[0][1])] + list(runs[1:])
        if bullet:
            p.space_after = Pt(4)
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.2))))
            pPr.set("indent", str(-int(Inches(0.2))))
            bu = etree.SubElement(pPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar")
            bu.set("char", "•")
        else:
            p.space_after = Pt(3)
        for t, fmt in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = font
            f.size = Pt(fmt.get("size", size))
            f.bold = fmt.get("bold", bold)
            f.italic = fmt.get("italic", italic)
            f.color.rgb = rgb(fmt.get("color", color))
    return box


def rect(slide, x: float, y: float, w: float, h: float, fill: str, edge: str | None = None, radius: float = 0.1,
         lw: float = 1.0, dashed: bool = False, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name: str | None = None):  # noqa: ANN001, ANN201
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / max(min(w, h), 0.01))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if edge:
        s.line.color.rgb = rgb(edge)
        s.line.width = Pt(lw)
        if dashed:
            s.line.dash_style = MSO_LINE.DASH
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def node(slide, x: float, y: float, w: float, h: float, title: str, sub: str | None = None, kind: str = "py",
         size: float = 12, sub_size: float | None = None, name: str | None = None):  # noqa: ANN201
    fill, edge = KINDS[kind]
    s = rect(slide, x, y, w, h, fill, edge, lw=1.25, name=name or f"node: {title}")
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.name, r.font.size, r.font.bold = FONT, Pt(size), True
    r.font.color.rgb = rgb(INK)
    if sub:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = sub
        r2.font.name, r2.font.size = FONT, Pt(sub_size or max(size - 2.5, 8.5))
        r2.font.color.rgb = rgb(MUTED)
    return s


def arrow(slide, x1: float, y1: float, x2: float, y2: float, color: str = LINE, lw: float = 1.75, dashed: bool = False,
          head: bool = True):  # noqa: ANN201
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    if dashed:
        c.line.dash_style = MSO_LINE.DASH
    if not head:
        return c
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def picture(slide, path: Path, x: float, y: float, w: float, h: float, alt: str, crop: tuple[float, float, float, float] = (0, 0, 0, 0),
            border: bool = True):  # noqa: ANN201
    """Places an image inside the (x, y, w, h) box without distortion (fit, centred). crop = l, t, r, b fractions."""
    with Image.open(path) as im:
        iw, ih = im.size
    iw_eff, ih_eff = iw * (1 - crop[0] - crop[2]), ih * (1 - crop[1] - crop[3])
    scale = min(w / iw_eff, h / ih_eff)
    pw, ph = iw_eff * scale, ih_eff * scale
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    pic.name = f"picture: {Path(path).stem}"
    if border:
        pic.line.color.rgb = rgb("CBD5E1")
        pic.line.width = Pt(0.75)
    return pic, (px, py, pw, ph)


def chip(slide, x: float, y: float, w: float, label: str, kind: str, size: float = 10.5, h: float = 0.3):  # noqa: ANN201
    fill, edge = {"V": ("DCFCE7", "14532D"), "U": ("FFEDD5", "7C2D12"), "S": ("EDE9FE", "4C1D95"),
                  "N": ("FEE2E2", "7F1D1D")}[kind]
    s = rect(slide, x, y, w, h, fill, edge, radius=h / 2, lw=1)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size, r.font.bold = FONT, Pt(size), True
    r.font.color.rgb = rgb(edge)
    return s


def table(slide, x: float, y: float, w: float, rows: list[list[str]], col_w: list[float], size: float = 11,
          row_h: float = 0.36, header_fill: str = NAVY, name: str = "table"):  # noqa: ANN201
    n, m = len(rows), len(rows[0])
    gs = slide.shapes.add_table(n, m, Inches(x), Inches(y), Inches(w), Inches(row_h * n))
    gs.name = name
    tbl = gs.table
    tblPr = tbl._tbl.tblPr
    style = tblPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId")
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid"
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = val
            r.font.name, r.font.size = FONT, Pt(size)
            if i == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(header_fill)
                r.font.bold = True
                r.font.color.rgb = rgb("FFFFFF")
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb("F5F8FC" if i % 2 == 0 else "FFFFFF")
                r.font.color.rgb = rgb(INK)
    return gs


def line_chart(slide, x: float, y: float, w: float, h: float, categories: list[str], series: dict[str, list[float]],
               colors: list[str], title: str, y_title: str = "", size: float = 10, legend: bool = True,
               smooth: bool = False):  # noqa: ANN201
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

    cd = CategoryChartData()
    cd.categories = categories
    for name, vals in series.items():
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = f"chart: {title}"
    ch = gf.chart
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    tp = ch.chart_title.text_frame.paragraphs[0]
    tp.runs[0].font.size, tp.runs[0].font.bold, tp.runs[0].font.name = Pt(size + 1), True, FONT
    tp.runs[0].font.color.rgb = rgb(NAVY)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size, ch.legend.font.name = Pt(size), FONT
    ch.font.size, ch.font.name = Pt(size), FONT
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("E2E8F0")
    va.format.line.color.rgb = rgb("CBD5E1")
    va.tick_labels.font.size = Pt(size)
    if y_title:
        va.has_title = True
        va.axis_title.text_frame.text = y_title
        va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    ch.category_axis.tick_labels.font.size = Pt(size)
    ch.category_axis.format.line.color.rgb = rgb("CBD5E1")
    for s, col in zip(ch.plots[0].series, colors):
        s.format.line.color.rgb = rgb(col)
        s.format.line.width = Pt(2.25)
        s.smooth = smooth
        s.marker.format.fill.solid()
        s.marker.format.fill.fore_color.rgb = rgb(col)
        s.marker.format.line.color.rgb = rgb(col)
        s.marker.size = 6
    return gf


def bar_chart(slide, x: float, y: float, w: float, h: float, categories: list[str], series: dict[str, list[float]],
              colors: list[str], title: str, stacked: bool = False, horizontal: bool = False, size: float = 10,
              legend: bool = True, number_format: str = "0.000"):  # noqa: ANN201
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

    cd = CategoryChartData()
    cd.categories = categories
    for name, vals in series.items():
        cd.add_series(name, vals)
    kind = {(False, False): XL_CHART_TYPE.COLUMN_CLUSTERED, (True, False): XL_CHART_TYPE.COLUMN_STACKED,
            (False, True): XL_CHART_TYPE.BAR_CLUSTERED, (True, True): XL_CHART_TYPE.BAR_STACKED}[(stacked, horizontal)]
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = f"chart: {title}"
    ch = gf.chart
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    tp = ch.chart_title.text_frame.paragraphs[0]
    tp.runs[0].font.size, tp.runs[0].font.bold, tp.runs[0].font.name = Pt(size + 1), True, FONT
    tp.runs[0].font.color.rgb = rgb(NAVY)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size)
    ch.font.size, ch.font.name = Pt(size), FONT
    ch.value_axis.major_gridlines.format.line.color.rgb = rgb("E2E8F0")
    ch.value_axis.tick_labels.font.size = Pt(size)
    ch.category_axis.tick_labels.font.size = Pt(size)
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(size - 1)
    plot.data_labels.number_format = number_format
    plot.data_labels.number_format_is_linked = False
    for s, col in zip(plot.series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(col)
    return gf


def header(slide, kicker: str, title: str, n: int, total: int) -> None:
    rect(slide, 0, 0, SW, 0.12, NAVY, None, shape=MSO_SHAPE.RECTANGLE, name="top bar")
    text(slide, 0.55, 0.3, 9, 0.3, kicker.upper(), size=11, color=BLUE, bold=True, name="kicker")
    text(slide, 0.55, 0.58, 12.2, 0.8, title, size=26, color=NAVY, bold=True, anchor=MSO_ANCHOR.TOP, name="title",
         line_spacing=0.95)
    text(slide, 0.55, 7.08, 9.5, 0.3, "Aegis-Net technical walkthrough - all data is synthetic; prototype, not production",
         size=9.5, color=MUTED, name="footer")
    text(slide, 11.4, 7.08, 1.4, 0.3, f"{n} / {total}", size=9.5, color=MUTED, align=PP_ALIGN.RIGHT, name="slide number")


def source_note(slide, s: str, y: float = 6.78) -> None:
    text(slide, 0.55, y, 12.2, 0.28, s, size=9.5, color=MUTED, italic=True, name="source note")
